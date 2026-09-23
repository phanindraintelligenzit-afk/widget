from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
import os

prs = Presentation('Slide_deck_planning_DPI_V-2.0_Final.pptx')

# 1. Title Slide
slide1 = prs.slides[0]
for shape in slide1.shapes:
    if shape.has_text_frame:
        if "deployment" in shape.text.lower():
            shape.text = "From agent deployment to performance optimization, efficiency gains, trust, and accountable autonomy"
        if "speedometer" in shape.text.lower():
            p = shape.text_frame.add_paragraph()
            p.text = "Digital Performance Index (DPI)"

# 2. AI Cost Paradox
slide2 = prs.slides[1]
for shape in slide2.shapes:
    if shape.has_text_frame:
        if "280x" in shape.text:
            shape.text = shape.text.replace("280x", "▼ 280x")
        if "320" in shape.text:
            shape.text = shape.text.replace("320%", "▲ 320%")
            shape.text = shape.text.replace("320x", "▲ 320x")
        if "cost of intelligence" in shape.text.lower():
            shape.text = "It is no longer the cost of intelligence, it is about coordinating intelligence for a meaningful scale."

# 3. The Elephant Slide (Slide 3)
slide3 = prs.slides[2]
# Insert the elephant image
img_path = r'C:\Users\User\.gemini\antigravity\brain\92a8bc58-10d4-4ff6-b9dc-6e7f48d6f5a2\.user_uploaded\media_1790077508309.jpg'
try:
    slide3.shapes.add_picture(img_path, Inches(1), Inches(2), width=Inches(5))
except:
    pass

for shape in slide3.shapes:
    if shape.has_text_frame:
        if "telemetry" in shape.text.lower() or "activities" in shape.text.lower():
            shape.text = "Engineering sees latency, Security sees risks, Finance sees costs. Like blind men defining an elephant, nobody sees the full picture."

# 4. Agent Evaluation Triangle
slide4 = prs.slides[3]
try: slide4.shapes.title.text = "Agent ecosystem needs to be observed, measured, and scaled—not just on token cost, but holistically across multiple dimensions."
except: pass

# 5. DPI Engine Slide (7 Dimensions)
slide5 = prs.slides[4]
try: slide5.shapes.title.text = "A performance score for every enterprise agent"
except: pass
for shape in slide5.shapes:
    if shape.has_text_frame:
        if "telemetry" in shape.text.lower():
            shape.text = "A performance score for every enterprise agent"
        if "P, Q" in shape.text:
            shape.text = "DPI = P, Q, E, G, R, C, V"

# 6. Enterprise Agent Control Tower
slide6 = prs.slides[5]
try: slide6.shapes.title.text = "THE ENTERPRISE AGENT CONTROL TOWER - ILLUSTRATION"
except: pass

txBox = slide6.shapes.add_textbox(Inches(7), Inches(1), Inches(2), Inches(1))
txBox.text_frame.text = "4 of 128 agents"

# 7. Scoring Layer
slide7 = prs.slides[6]
try: slide7.shapes.title.text = "Scoring layer seamlessly integrates into any agent ecosystem"
except: pass
for shape in slide7.shapes:
    if shape.has_text_frame:
        if "not another" in shape.text.lower() or "not yet another" in shape.text.lower():
            shape.text = "An interpretation layer for the evidence, not yet another observability stack."

# 8. Performance Contract
slide8 = prs.slides[7]
try: slide8.shapes.title.text = "Advancing Autonomy: From L1 to L4"
except: pass
for shape in slide8.shapes:
    if shape.has_text_frame:
        shape.text = "Measuring agents allows us to move them from L1 to L4 autonomy, put them on a PIP, or remove the human-in-the-loop. We move from individual metrics to team and department aggregation."

# 9. Closing Slide
slide9 = prs.slides[-1]
for shape in slide9.shapes:
    if shape.has_text_frame:
        if "Thank You" in shape.text or "thank you" in shape.text.lower():
            shape.text = "Feel free to reach out to us for a demo of DPI in action"

prs.save('DPI- Updated 2026_FINAL_VERSION.pptx')
print("Complete PPT generation successful!")
