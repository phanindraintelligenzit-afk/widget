from pptx import Presentation

ppt_path = "Slide_deck_planning_ DPI V-2.0.pptx"
try:
    prs = Presentation(ppt_path)
    with open("ppt_dump.txt", "w", encoding="utf-8") as f:
        for i, slide in enumerate(prs.slides):
            f.write(f"--- Slide {i+1} ---\n")
            for shape in slide.shapes:
                if hasattr(shape, "text"):
                    f.write(shape.text + "\n")
except Exception as e:
    print(f"Error: {e}")
