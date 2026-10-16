from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
import os

def apply_style(slide, title_text, body_text_list, notes_text=""):
    # Set background color to dark blue
    background = slide.background
    fill = background.fill
    fill.solid()
    fill.fore_color.rgb = RGBColor(11, 17, 32) # Dark Blue
    
    # Title formatting
    title_shape = slide.shapes.title
    title_shape.text = title_text
    title_frame = title_shape.text_frame
    for paragraph in title_frame.paragraphs:
        for run in paragraph.runs:
            run.font.color.rgb = RGBColor(255, 204, 0) # Intelligenz Yellow
            run.font.bold = True
            
    # Body formatting
    body_shape = slide.placeholders[1]
    tf = body_shape.text_frame
    tf.text = body_text_list[0]
    tf.paragraphs[0].font.color.rgb = RGBColor(255, 255, 255) # White
    tf.paragraphs[0].font.size = Pt(28)
    
    for text_item in body_text_list[1:]:
        p = tf.add_paragraph()
        p.text = text_item
        p.font.color.rgb = RGBColor(255, 255, 255) # White
        p.font.size = Pt(24)
        
    # Speaker Notes
    if notes_text:
        notes_slide = slide.notes_slide
        notes_slide.notes_text_frame.text = notes_text

def create_styled_ppt():
    prs = Presentation()
    
    # Slide 1: Title Slide
    slide = prs.slides.add_slide(prs.slide_layouts[0])
    
    # Background
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = RGBColor(11, 17, 32)
    
    # Title
    title = slide.shapes.title
    title.text = "DPI Germany-India"
    for run in title.text_frame.paragraphs[0].runs:
        run.font.color.rgb = RGBColor(255, 204, 0)
        run.font.bold = True
        run.font.size = Pt(54)
        
    # Subtitle
    subtitle = slide.placeholders[1]
    subtitle.text = "The Complete AI Measurement Engine"
    for run in subtitle.text_frame.paragraphs[0].runs:
        run.font.color.rgb = RGBColor(255, 255, 255)
        run.font.size = Pt(32)
    
    # Slide 2: Introduction
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    apply_style(slide, 
                "1. Introduction (The Elephant Concept)", 
                ["Moving beyond 'Can it do the work?' to 'Can we trust and scale it?'",
                 "Like the Blind Men and the Elephant, measuring only one part of an AI misses the big picture.",
                 "DPI Germany-India sees the whole AI."],
                "Respected everyone, good afternoon. Today I am going to present DPI. First, I will explain what DPI is and why we built it. DPI stands for Digital Performance Index. It is a measurement engine that measures AI-agent performance.\n\nLike the story of the blind men and the elephant, if we only measure one thing, we miss the big picture. DPI helps us understand productivity, quality, risk, governance, validation, and cost as a whole before we trust and scale an AI agent.")

    # Slide 3: Core Concept
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    apply_style(slide,
                "2. How DPI Works",
                ["The Core Workflow:",
                 "AI Agent -> Connect -> Measure -> Score -> Understand -> Improve -> Trust -> Scale"],
                "How does DPI work? The core workflow is simple and continuous: AI Agent connects to the DPI Engine. We measure the data, calculate a score, and help you understand the performance. This allows you to improve the agent, build trust, and ultimately scale it.")

    # Slide 4: Dimensions
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    apply_style(slide,
                "3. What Does DPI Measure? (The Paradox)",
                ["Expanding the traditional Speed/Quality/Cost paradox into 7 Dimensions:",
                 "1. Productivity (P)", "2. Quality (Q)", "3. Execution (E)", "4. Governance (G)", "5. Risk (R)", "6. Validation (V)", "7. Cost (C)"],
                "Traditionally in engineering, there is a paradox: you have to choose between speed, quality, and cost. DPI expands this triangle into seven dimensions to measure the complete reality. We measure: Productivity, Quality, Execution, Governance, Risk, Validation, and Cost.")

    # Slide 5: UI Flow Overview
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    apply_style(slide,
                "4. Application Demo (The Complete Flow)",
                ["Complete UI Flow:",
                 "Login -> Onboard Agent -> Configuration -> Dashboard -> Rating -> Profile -> Resources -> Sign Out"],
                "Now, let me explain how the application works in practice. The complete flow is: Login, Onboard Agent, Configuration, Dashboard, Rating, Profile, Resources, and Sign Out.")

    # Slide 6: Walkthrough
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    apply_style(slide,
                "5. Step-by-Step Walkthrough",
                ["Login: Secure access to the 7 main sections.",
                 "Onboard: Enter details & Submit.",
                 "Configuration: Set the 7 weightages (Must equal 100%).",
                 "Dashboard: See real-time scores and profiles.",
                 "Rating & Profile: Manual reviews and historical tracking."],
                "Step 1, the user logs in. Step 2, they go to Onboard Agent and click Submit. Step 3, they are redirected to Configuration to update the 7 weightages, which must equal 100%. After saving, step 4, the calculated scores reflect on the Dashboard. Then the user can provide ratings which update the Profile section.")

    # Slide 7: Resources
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    apply_style(slide,
                "6. Integrations & Resources",
                ["18 Available Resources Integrated:",
                 "Execution: Langfuse, OpenTelemetry, Jaeger",
                 "Quality: Ragas, DeepEval, AgentOps",
                 "Risk: Falco, Sentry, Prometheus",
                 "Governance: OpenMetadata, OPA, Keycloak",
                 "Cost: Langfuse, OpenCost, OpenLIT",
                 "Productivity & Validation included"],
                "The Resources section contains 18 available resources. DPI is designed to work with enterprise tools already generating data. We use Langfuse and OpenTelemetry for Execution, Ragas and DeepEval for Quality, Falco for Risk, and so on.")

    # Slide 8: Value
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    apply_style(slide,
                "7. Value & Conclusion",
                ["Measure -> Trust -> Scale",
                 "Unified view for Business, Engineering, Risk, and Compliance.",
                 "Evidence-backed dimension scoring."],
                "Finally, why should a customer trust our score? Because the DPI score is not a black-box number. We break the score down into seven dimensions and trace them back to exact evidence. The ultimate value is that it brings Business, Engineering, Risk, and Compliance into one unified view. Measure, Trust, Scale. Thank you.")

    prs.save(r'D:\DPI-LS\widget\DPI Germany-india.pptx')

if __name__ == '__main__':
    create_styled_ppt()
