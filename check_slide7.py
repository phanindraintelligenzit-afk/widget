from pptx import Presentation
prs = Presentation("Slide_deck_planning_DPI_V-2.0_Updated.pptx")
slide = prs.slides[6]
for shape in slide.shapes:
    if hasattr(shape, "text"):
        print(repr(shape.text))
