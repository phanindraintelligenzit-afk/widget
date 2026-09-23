from pptx import Presentation

ppt_path = "Slide_deck_planning_DPI_V-2.0_Updated.pptx"
out_path = "Slide_deck_planning_DPI_V-2.0_Final.pptx"
prs = Presentation(ppt_path)
slide = prs.slides[6] # Slide 7

changed = False
for shape in slide.shapes:
    if not hasattr(shape, "text"): continue
    
    if "API Interface" in shape.text or "Telemetry + API + MCP" in shape.text:
        # Check if the text length matches the box text reasonably
        if len(shape.text) < 50:
            shape.text = "API Interface (OpenTelemetry, REST APIs, and MCP integrations)"
            changed = True

if changed:
    prs.save(out_path)
    print("Successfully created Final PPT!")
else:
    print("Could not find the target text box on Slide 7.")
