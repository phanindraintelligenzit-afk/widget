from pptx import Presentation

prs = Presentation('Slide_deck_planning_DPI_V-2.0_Final.pptx')

# Slide 6 is index 5
slide6 = prs.slides[5]

for shape in slide6.shapes:
    if hasattr(shape, "text_frame") and shape.text_frame:
        tf = shape.text_frame
        # Replace the terms
        if "1. Science" in tf.text:
            tf.text = tf.text.replace("1. Science", "1. Science: Competency")
            tf.text = tf.text.replace("2. Art", "2. Art: Collaboration")
            tf.text = tf.text.replace("3. Security", "3. Security: Maturity")
            tf.text = tf.text.replace("Science, Art, and Security", "Science (Competency), Art (Collaboration), and Security (Maturity)")

prs.save('Slide_deck_planning_DPI_V-2.0_Final.pptx')
print("Slide 6 Updated!")
