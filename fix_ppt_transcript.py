from pptx import Presentation
from pptx.util import Inches, Pt
import os

prs = Presentation('Slide_deck_planning_DPI_V-2.0_Final.pptx')

# We will create a new presentation for the output
# But to avoid breaking layouts, let's just do text replacements.

for slide in prs.slides:
    for shape in slide.shapes:
        if not shape.has_text_frame:
            continue
        
        # Slide 1:
        if "From agent deployment to performance" in shape.text:
            shape.text = "From agent deployment to performance optimization, efficiency gains, trust, and accountable autonomy."
            
        # Slide 2:
        if "280x" in shape.text:
            shape.text = shape.text.replace("280x", "▼ 280x")
        if "320x" in shape.text:
            shape.text = shape.text.replace("320x", "▲ 320x")
            
        # Architecture Slide
        if "WHERE THE SCORING LAYER SITS" in shape.text or "Where the scoring layer sits" in shape.text.lower():
            shape.text = "THE SCORING LAYER SEAMLESSLY INTEGRATES INTO ANY AGENT ECOSYSTEM"
        if "not another observability stack" in shape.text.lower() or "not yet another" in shape.text.lower():
            shape.text = shape.text.replace(shape.text, "It is an interpretation layer for the evidence.")

        # Dashboard Slide
        if "THE ENTERPRISE AGENT CONTROL TOWER" in shape.text:
            if "ILLUSTRATION" not in shape.text:
                shape.text = shape.text.replace("THE ENTERPRISE AGENT CONTROL TOWER", "THE ENTERPRISE AGENT CONTROL TOWER (ILLUSTRATION)")

        # Removing Appendix text
        if "Appendix -" in shape.text:
            shape.text = shape.text.replace("Appendix - ", "")

prs.save('DPI- Updated 2026.pptx')
print("Text updates complete!")
