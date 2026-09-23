from pptx import Presentation
from pptx.util import Inches, Pt
import os

ppt_path = 'Slide_deck_planning_DPI_V-2.0_Final.pptx'
prs = Presentation(ppt_path)

# Slide 3 (index 2): Aspects on the Spent and Token Consumption increase
slide3 = prs.slides[2]
txBox = slide3.shapes.add_textbox(Inches(0.5), Inches(5), Inches(9), Inches(2))
tf = txBox.text_frame
p = tf.add_paragraph()
p.text = "Additional Aspects:"
p.font.bold = True
p.font.size = Pt(20)
p2 = tf.add_paragraph()
p2.text = "1. Spend has increased by 40%+"
p2.font.size = Pt(18)
p3 = tf.add_paragraph()
p3.text = "2. Must measure Efficiency / ROI at the Logical unit"
p3.font.size = Pt(18)
p4 = tf.add_paragraph()
p4.text = "3. Comparison strategy: Agentic Workflows vs Traditional Execution"
p4.font.size = Pt(18)

# Move 8th slide (index 7) below the 3rd slide (index 2) -> so it becomes index 3
xml_slides = prs.slides._sldIdLst
slides = list(xml_slides)
slide8 = slides[7]
xml_slides.remove(slide8)
xml_slides.insert(3, slide8)

# Add Slide: Why is the above slide info important (Pfizer ROI)
layout = prs.slide_layouts[0]
slide_pfizer = prs.slides.add_slide(layout)
try: slide_pfizer.shapes.title.text = "Why Measuring Agent ROI is Critical"
except: pass
txBox2 = slide_pfizer.shapes.add_textbox(Inches(1), Inches(1.5), Inches(8), Inches(4))
tf2 = txBox2.text_frame
p = tf2.add_paragraph()
p.text = "Why is this important?"
p.font.bold = True
p.font.size = Pt(24)
p2 = tf2.add_paragraph()
p2.text = "Without measuring ROI at the logical unit, AI scales as a cost-center rather than a profit-driver."
p2.font.size = Pt(20)
p3 = tf2.add_paragraph()
p3.text = "\nExample of benefits of measuring ROI:"
p3.font.bold = True
p3.font.size = Pt(24)
p4 = tf2.add_paragraph()
p4.text = "Top pharma companies like Pfizer measure the exact ROI of their AI systems by tracking time-saved in drug discovery data analysis and regulatory compliance checks against the actual token cost."
p4.font.size = Pt(20)

# Move Pfizer slide to be right after the moved 8th slide (index 4)
slides = list(xml_slides)
pfizer_xml = slides[-1]
xml_slides.remove(pfizer_xml)
xml_slides.insert(4, pfizer_xml)

# Add Slide: Science -> Art -> Security
slide_sas = prs.slides.add_slide(layout)
try: slide_sas.shapes.title.text = "The Evolution of AI Measurement"
except: pass
txBox3 = slide_sas.shapes.add_textbox(Inches(1), Inches(1.5), Inches(8), Inches(4))
tf3 = txBox3.text_frame
p = tf3.add_paragraph()
p.text = "1. Science"
p.font.bold = True
p.font.size = Pt(24)
p2 = tf3.add_paragraph()
p2.text = "The foundational engineering, LLMs, and compute."
p2.font.size = Pt(20)
p3 = tf3.add_paragraph()
p3.text = "\n2. Art"
p3.font.bold = True
p3.font.size = Pt(24)
p4 = tf3.add_paragraph()
p4.text = "Prompting, agentic behaviors, and workflows."
p4.font.size = Pt(20)
p5 = tf3.add_paragraph()
p5.text = "\n3. Security"
p5.font.bold = True
p5.font.size = Pt(24)
p6 = tf3.add_paragraph()
p6.text = "Governance, risk, and compliance."
p6.font.size = Pt(20)
p7 = tf3.add_paragraph()
p7.text = "\n--> DPI (Formula) - Measuring the balance of all three."
p7.font.bold = True
p7.font.size = Pt(24)

# Move SAS slide to index 5
slides = list(xml_slides)
sas_xml = slides[-1]
xml_slides.remove(sas_xml)
xml_slides.insert(5, sas_xml)

prs.save('Slide_deck_planning_DPI_V-2.0_Final.pptx')
print("Successfully overwrote the V-2.0 file!")
